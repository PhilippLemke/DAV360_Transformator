# Regressionstests: der Transformator muss für jede Typ/Variante-Kombination
# exakt die hinterlegten Beispiel-Exporte in tests/fixtures/expected/ erzeugen.
# Damit fällt sofort auf, wenn eine Änderung an profiles.yaml/mapping.yaml/
# transform_functions.py das Ergebnis unbeabsichtigt verändert.

import os

import openpyxl
import pytest

import transformator

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures", "expected")

CASES = [
    ("touren", "tak", "dummy_daten/touren_tak_beispiel.xlsx", "touren_tak.xlsx"),
    ("touren", "tr", "dummy_daten/touren_tr_beispiel.xlsx", "touren_tr.xlsx"),
    ("veranstaltungen", "tr", "dummy_daten/veranstaltungen_tr_beispiel.xlsx", "veranstaltungen_tr.xlsx"),
    ("touren", "gruppen", "dummy_daten/gruppen_beispiel.xlsx", "gruppen_touren.xlsx"),
    ("veranstaltungen", "gruppen", "dummy_daten/gruppen_beispiel.xlsx", "gruppen_veranstaltungen.xlsx"),
]


@pytest.fixture(scope="module")
def loaded_config():
    os.chdir(REPO_ROOT)
    config = transformator.load_yaml("config.yaml")
    profiles = transformator.load_yaml(config["profiles_file"])
    mapping = transformator.Mapping(transformator.load_yaml(config["mapping_file"]))
    return config, profiles, mapping


def _sheet_values(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    return [
        [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        for r in range(1, ws.max_row + 1)
    ]


@pytest.mark.parametrize("typ,variante,input_file,fixture_name", CASES)
def test_output_matches_fixture(loaded_config, typ, variante, input_file, fixture_name):
    config, profiles, mapping = loaded_config
    out_file = transformator.run(typ, variante, input_file, config, profiles, mapping)
    try:
        actual = _sheet_values(out_file)
        expected = _sheet_values(os.path.join(FIXTURES, fixture_name))
        assert actual == expected
    finally:
        if os.path.exists(out_file):
            os.remove(out_file)


def test_kurse_not_implemented(loaded_config, capsys):
    config, profiles, mapping = loaded_config
    with pytest.raises(SystemExit):
        transformator.run("kurse", "irgendeine", "egal.xlsx", config, profiles, mapping)
    assert "noch nicht implementiert" in capsys.readouterr().out
