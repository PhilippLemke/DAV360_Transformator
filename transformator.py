#!/usr/bin/env python3
# DAV360 Transformator
# GNU General Public License v3.0
#
# Einheitlicher Transformator für DAV360/Pimcore-Importe. Ersetzt die
# früheren Einzelskripte (TourenTransformatorMSF.py, TourenTransformatorTR.py,
# VeranstaltungsTransformatorMSF.py, GruppenTransformatorMSF.py).
#
# Aufruf:
#   python transformator.py <typ> --variante <variante> <eingabedatei.xlsx>
#
# Konfiguration:
#   config.yaml    - Tool-Verhalten (Pfade, Season-Logik, Output-Ordner)
#   mapping.yaml   - Sektionsspezifische Stammdaten (Kategorien, Tourenführer, Gruppen, ...)
#   profiles.yaml  - Spalten-Zuordnung je Typ/Variante

import argparse
import os
import sys
from dataclasses import dataclass
from datetime import datetime

import openpyxl
import pandas
import yaml

import transform_functions as tf


def load_yaml(path):
    if not os.path.exists(path):
        print(f"ERROR: Konfigurationsdatei {path} fehlt! -> Abbruch")
        sys.exit(1)
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


class Mapping:
    """Sektionsspezifische Stammdaten aus mapping.yaml (ersetzt Keys.xlsx)."""

    def __init__(self, data):
        self.kategorie = data.get("kategorie", {})
        self.technik = data.get("technik", {})
        self.ausdauer = data.get("ausdauer", {})
        self.saison = data.get("saison", {})
        self.eventart = data.get("eventart", {})
        self.klassifizierung = data.get("klassifizierung", {})
        self.kursstufe = data.get("kursstufe", {})
        self.tourenfuehrer = data.get("tourenfuehrer", {})
        self.gruppen = data.get("gruppen", {})
        self.gruppe_kategorie = data.get("gruppe_kategorie", {})


@dataclass
class SeasonContext:
    season_name: str  # "Sommer" oder "Winter"
    program_year: int
    season_id: str  # z.B. "Sommer2026"


def compute_season(config, today=None):
    today = today or datetime.today()
    winter_after = config.get("season", {}).get("winter_starts_after_month", 6)
    if today.month > winter_after:
        season_name = "Winter"
        program_year = today.year + 1
    else:
        season_name = "Sommer"
        program_year = today.year
    return SeasonContext(season_name, program_year, f"{season_name}{program_year}")


def resolve_output_path(config, typ, variante, season_ctx):
    template = config["output"]["filenames"].get(typ, {}).get(variante)
    if template is None:
        template = f"DAV_{typ}export_{{season_id}}.xlsx"
    filename = template.format(season_id=season_ctx.season_id)
    out_dir = config.get("output", {}).get("dir", "export")
    out_dir = out_dir if os.path.isdir(out_dir) else "."
    return os.path.join(out_dir, filename)


def row_matches_filter(row, variant_spec):
    row_filter = variant_spec.get("filter")
    if row_filter is None:
        return True
    return str(row.get(row_filter["column"])) == row_filter["equals"]


def check_required_columns(input_file, typ, variante, data_in, variant_spec):
    required = tf.required_columns_for_variant(variant_spec)
    missing = sorted(col for col in required if col not in data_in.columns)
    if not missing:
        return
    print(f"ERROR: Eingabedatei '{input_file}' passt nicht zu Typ '{typ}' / Variante '{variante}'.")
    print(f"Es fehlen folgende erwartete Spalten: {', '.join(missing)}")
    print(f"Vorhandene Spalten: {', '.join(str(c) for c in data_in.columns)}")
    print("Bitte prüfen, ob die richtige Eingabedatei bzw. die richtige --variante gewählt wurde.")
    sys.exit(1)


def transform_input(input_file, typ, variante, typ_spec, variant_spec, mapping, season_ctx):
    data_in = pandas.read_excel(input_file).reset_index()
    check_required_columns(input_file, typ, variante, data_in, variant_spec)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None  # ein frisch erzeugtes Workbook hat immer ein aktives Blatt
    sheet.title = typ_spec.get("sheet_name", typ_spec)

    output_columns = typ_spec["output_columns"]
    for column_index, column_name in enumerate(output_columns, start=1):
        sheet.cell(row=1, column=column_index).value = column_name

    fields = variant_spec.get("fields", {})
    out_row = 2
    for _, row in data_in.iterrows():
        if not row_matches_filter(row, variant_spec):
            continue
        row_label = row.get("Bezeichnung/Titel", row.get("Titel", ""))
        print(f"Verarbeite Zeile: {row_label}")
        for column_index, column_name in enumerate(output_columns, start=1):
            spec = fields.get(column_name)
            if spec is None:
                continue
            try:
                value = tf.dispatch(spec, row, mapping, season_ctx)
            except Exception as exc:
                print(f"ERROR: Zeile '{row_label}', Ausgabefeld '{column_name}' (transform: {spec.get('transform')}): {exc}")
                print("Vermutlich fehlt in dieser Zeile ein benötigter Eingabewert (z.B. ein Datum). Bitte Eingabedatei prüfen.")
                sys.exit(1)
            sheet.cell(row=out_row, column=column_index).value = value  # type: ignore[union-attr]
        out_row += 1

    return workbook


def run(typ, variante, input_file, config, profiles, mapping):
    typ_spec = profiles.get(typ)
    if typ_spec is None:
        print(f"ERROR: Unbekannter Typ '{typ}'. Verfügbar: {', '.join(profiles.keys())}")
        sys.exit(1)

    variants = typ_spec.get("variants", {})
    if not variants:
        print(f"'{typ}' ist als Typ vorgesehen, aber noch nicht implementiert (keine Variante in profiles.yaml).")
        sys.exit(1)

    variant_spec = variants.get(variante)
    if variant_spec is None:
        print(f"ERROR: Unbekannte Variante '{variante}' für Typ '{typ}'. Verfügbar: {', '.join(variants.keys())}")
        sys.exit(1)

    if not os.path.exists(input_file):
        print(f"ERROR: Eingabedatei {input_file} existiert nicht! -> Abbruch")
        sys.exit(1)

    season_ctx = compute_season(config)
    print(f"Typ: {typ}, Variante: {variante}, Season: {season_ctx.season_id}")
    print(f"Verwende Eingabedatei: {input_file}")

    workbook = transform_input(input_file, typ, variante, typ_spec, variant_spec, mapping, season_ctx)

    out_file = resolve_output_path(config, typ, variante, season_ctx)
    workbook.save(out_file)
    print(f"Geschrieben: {out_file}")
    return out_file


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        prog="transformator.py",
        description="Konvertiert Touren-/Veranstaltungs-Eingabeformulare für den DAV360-Import.",
    )
    parser.add_argument("typ", help="z.B. touren, veranstaltungen, kurse")
    parser.add_argument("--variante", required=True, help="z.B. msf, tr, gruppen")
    parser.add_argument("eingabedatei", help="Pfad zur Eingabe-Exceldatei")
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    config = load_yaml("config.yaml")
    profiles = load_yaml(config.get("profiles_file", "profiles.yaml"))
    mapping = Mapping(load_yaml(config.get("mapping_file", "mapping.yaml")))
    run(args.typ, args.variante, args.eingabedatei, config, profiles, mapping)


if __name__ == "__main__":
    main()
